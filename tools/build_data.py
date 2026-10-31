# Converts the CSE workbook (v7) into data/curriculum.json, data/topic-projects.json and data/question-bank.json.
#   python3 tools/build_data.py <workbook.xlsx | workbook-export.txt>
# The .txt form is the tab-separated "## Sheet:" export. Some exports stop at 1,000 rows per sheet; when that happens
# this script keeps the earlier data for the missing sessions (marked dataSource "carried-over") and never invents rows.
# Stable app IDs are kept so saved progress still matches: modules M01.., projects PR01.., original questions P001..
import sys, json, re, os

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, '..', 'data')
src = sys.argv[1]


def load_sheets(path):
    sheets = {}
    if path.lower().endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            try:
                hdr = [str(h).strip() if h is not None else '' for h in next(it)]
            except StopIteration:
                continue
            out = []
            for r in it:
                if r[0] is None:
                    continue
                out.append({h: ('' if v is None else (v if isinstance(v, (int, float)) else str(v).strip())) for h, v in zip(hdr, r)})
            sheets[ws.title] = (hdr, out)
        return sheets
    cur, lines = None, {}
    for line in open(path, encoding='utf-8').read().split('\n'):
        if line.startswith('## Sheet: '):
            cur = line[10:].strip(); lines[cur] = []; continue
        if cur is not None and line.strip() and not line.startswith('... (truncated'):
            lines[cur].append(line)
    for name, ls in lines.items():
        if len(ls) < 2:
            sheets[name] = ([], []); continue
        hdr = [h.strip() for h in ls[0].split('\t')]
        rows = []
        for l in ls[1:]:
            c = l.split('\t')
            if len(c) < 2:
                continue
            c += [''] * (len(hdr) - len(c))
            rows.append({h: v.strip() for h, v in zip(hdr, c)})
        sheets[name] = (hdr, rows)
    return sheets


S = load_sheets(src)


def raw_rows(name):   # every non-empty row of a sheet as a list of non-empty cells (for free-form sheets such as Start Here)
    if src.lower().endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(src, read_only=True, data_only=False)
        return [[str(c).strip() for c in r if c is not None and str(c).strip()] for r in wb[name].iter_rows(values_only=True) if any(c is not None for c in r)]
    out, on = [], False
    for line in open(src, encoding='utf-8').read().split('\n'):
        if line.startswith('## Sheet: '): on = (line[10:].strip() == name); continue
        if on and line.strip() and not line.startswith('... (truncated'):
            cells = [c.strip() for c in line.split('\t') if c.strip()]
            if cells: out.append(cells)
    return out
def rows(name): return S[name][1] if name in S else []
def num(v, default=None):
    try: return int(float(v))
    except (TypeError, ValueError): return default
def txt(v): return None if v is None or str(v).strip() in ('', '-') else str(v).strip()
def links(s): return [u.strip() for u in s.split('|') if u.strip() and u.strip() != '-'] if s else []
def ids(s): return [x.strip() for x in str(s or '').split(',') if x.strip()]
def dayList(s): return [int(x) for x in re.findall(r'\d+', str(s or ''))]
def mid(x): return 'M' + x[3:] if x and x.startswith('MOD') else x          # MOD01 -> M01 (app id)
def pid(x): return 'PR' + x[3:] if x and x.startswith('PRJ') else x          # PRJ01 -> PR01 (app id)

old_path = os.path.join(DATA, 'curriculum.json')
old = json.load(open(old_path, encoding='utf-8')) if os.path.exists(old_path) else None

# ---------- modules
modules = []
for m in rows('Module Catalogue'):
    modules.append({'id': mid(m['Module ID']), 'catalogId': m['Module ID'], 'name': m['Module'], 'phase': m['Phase'], 'semester': m['Semester target'], 'domain': m['Domain'],
        'days': num(m['Learning days']), 'startDay': num(m['Start day']), 'endDay': num(m['End day']), 'outcome': m['Expected outcome'],
        'resource': m['Free resource'], 'resourceLinks': links(m['Resource URL']), 'extensionRule': m['Extension rule']})
modByName = {m['name']: m for m in modules}
total = sum(m['days'] for m in modules)

# ---------- topics (294)
topics = []
for t in rows('Topic Catalogue'):
    topics.append({'id': t['Topic ID'], 'moduleId': mid(t['Module ID']), 'name': t['Topic / concept'], 'days': dayList(t['Related session days']),
        'firstDay': num(t['First learning day']), 'lastDay': num(t['Last learning day']), 'outcome': t['Learning outcome (source)'], 'confidence': t['Mapping confidence'],
        'note': t['Definition / source note'], 'projectIds': [pid(x) for x in ids(t['Related project IDs'])]})
topicOfDay = {}
for t in topics:
    for d in t['days']: topicOfDay[d] = t
topicById = {t['id']: t for t in topics}
for m in modules: m['topicIds'] = [t['id'] for t in topics if t['moduleId'] == m['id']]

# ---------- projects
cat = rows('Project Catalogue'); ms = {r['Project']: r for r in rows('Project Milestones')}; tl = rows('Project Version Timeline')
meta = {r['Project']: r for r in rows('Project Metadata')}; mtm = {(r['Project ID'], r['Version']): r for r in rows('Milestone Target Map')}
oldProj = {p['name']: p for p in (old or {}).get('projects', [])}
projects = []
for i, c in enumerate(cat, 1):
    n = c['Project']; m = ms[n]; pm = meta.get(n, {}); wid = c['Project ID']; p_id = 'PR%02d' % i
    assert pid(wid) == p_id, (wid, p_id)
    vs = []
    for t in [t for t in tl if t['Project'] == n]:
        tgt = t['Target learning day / track']; vn = t['Version']; mm = mtm.get((wid, vn), {})
        isnum = str(tgt).strip().isdigit()
        vs.append({'id': '%s-%s' % (p_id, vn), 'name': vn, 'targetDay': int(tgt) if isnum else None, 'targetText': None if isnum else tgt,
            'targetType': mm.get('Target type'), 'startDay': num(mm.get('Mapped start day')), 'endDay': num(mm.get('Mapped end day')), 'topicIds': ids(mm.get('Related topic IDs')),
            'mappingStatus': mm.get('Mapping status'), 'unlocksAfter': t['Unlocks after'], 'changes': t['What changes in this version'], 'acceptance': t['Acceptance checks'],
            'folder': t['Canonical folder'], 'workbookStatus': t['Status']})
    fm = m['First planned milestone day']
    p = {'id': p_id, 'workbookId': wid, 'name': n, 'category': c['Category'], 'coreSkills': c['Core skills'], 'priority': c['Portfolio priority'],
        'workbookStatus': c['Status'], 'notes': c['Notes'], 'catalogueVersions': [c['Version 1'], c['Version 2'], c['Version 3'], c['Version 4 / extension']],
        'prerequisite': m['Prerequisite / introduction'], 'firstMilestoneDay': int(fm) if str(fm).strip().isdigit() else fm, 'versionRoadmap': m['Version roadmap'],
        'folder': m['Canonical folder'], 'publishableDefinition': m['Definition of a publishable milestone'], 'continuityRule': m['Continuity rule'], 'versions': vs,
        'purpose': pm.get('Purpose / problem solved'), 'stack': pm.get('Proposed technology stack'), 'unlockSkills': pm.get('Prerequisites / unlock skills'),
        'architecture': pm.get('Suggested architecture / components'), 'testing': pm.get('Testing / acceptance expectations'), 'githubDocs': pm.get('GitHub / documentation expectations'),
        'deployment': pm.get('Deployment / execution target'), 'moduleIds': [mid(x) for x in ids(pm.get('Related module IDs'))], 'topicIds': ids(pm.get('Related topic IDs')),
        'sessionDays': dayList(pm.get('Related session days')), 'reviewStatus': pm.get('Review status'), 'provenance': pm.get('Metadata provenance')}
    p['relatedDays'] = sorted({v['targetDay'] for v in vs if v['targetDay']} | ({p['firstMilestoneDay']} if isinstance(p['firstMilestoneDay'], int) else set()))
    projects.append(p)
projectTopics = [{'projectId': pid(r['Project ID']), 'topicId': r['Topic ID'], 'relationship': r['Relationship type'], 'note': r['Confidence / review note']} for r in rows('Project Topic Map')]

# ---------- sessions
expectedCols = ['Roadmap Day']
daily = {}
for r in rows('Daily Roadmap'):
    d = num(r['Roadmap Day'])
    if d: daily[d] = r
sessions = []; carried = []
for d in range(1, total + 1):
    mod = next(m for m in modules if m['startDay'] <= d <= m['endDay']); tp = topicOfDay.get(d)
    r = daily.get(d)
    if r:
        cp = r['Project / codebase checkpoint'] or ''; prj = None
        mt = re.match(r'^(.+?): checkpoint', cp) or re.match(r'^MILESTONE: (.+?) V\d', cp)
        byName = {p['name']: p for p in projects}
        if mt and mt.group(1) in byName: prj = byName[mt.group(1)]['id']
        sessions.append({'id': 'D%04d' % d, 'day': d, 'phase': r['Phase'], 'semester': r['Semester target'], 'domain': r['Domain'], 'moduleId': mid(r['Module ID']), 'module': r['Module'],
            'moduleDay': num(r['Module Day']), 'moduleTotal': num(r['Module Total Days']), 'title': r['Lesson / Focus'], 'outcome': r['Learning outcome'], 'task': r['Detailed task for this session'],
            'practice': r['Practice task (attempt yourself)'], 'practiceType': r['Practice type'], 'platform': r['Platform guidance'], 'problemLinks': [u for u in (r['Problem 1 direct link'], r['Problem 2 direct link']) if u and u != '-'],
            'resource': r['Free resource'], 'resourceLinks': links(r['Resource link']), 'checkpoint': r['Project / codebase checkpoint'], 'projectAction': r['Project action'], 'folder': r['Folder / file guidance'],
            'github': r['GitHub action'], 'time': r['Suggested time'], 'revision': r['Revision instruction'], 'mastery': r['Mastery evidence / exit check'],
            'workbookStatus': txt(r['Status']) if 'Status' in r else None, 'workbookRemarks': txt(r['Personal remarks / blockers']), 'projectId': prj,
            'topicId': r['Topic ID'], 'topic': r['Topic name'], 'relatedProjectIds': [pid(x) for x in ids(r['Related project IDs'])], 'questionIds': ids(r['Related practice question IDs']), 'dataSource': 'workbook-v7'})
    else:
        o = dict(old['sessions'][d - 1]) if old and len(old['sessions']) >= d else {'id': 'D%04d' % d, 'day': d}
        o.update({'moduleId': mod['id'], 'topicId': tp['id'] if tp else None, 'topic': tp['name'] if tp else None, 'relatedProjectIds': o.get('relatedProjectIds', []),
            'questionIds': [], 'problemLinks': o.get('problemLinks', []), 'dataSource': 'carried-over'})
        sessions.append(o); carried.append(d)
assert [s['day'] for s in sessions] == list(range(1, total + 1))

# ---------- practice questions: merge workbook questions into the existing bank (never drop existing ids)
bankPath = os.path.join(DATA, 'question-bank.json')
bank = json.load(open(bankPath, encoding='utf-8')) if os.path.exists(bankPath) else {'meta': {}, 'questions': []}
qmap = {q['id']: q for q in bank['questions']}
pqm = {r['Question ID']: r for r in rows('Practice Question Map')}
dayOfQ = {}
for s in sessions:
    for q in s['questionIds']: dayOfQ.setdefault(q, []).append(s['day'])
added = updated = 0
for r in rows('Practice Bank'):
    qid = r['ID']; mp = pqm.get(qid, {}); days = dayList(r.get('Related session days')) or dayOfQ.get(qid, [])
    tid = r.get('Mapped Topic ID') or None; mo = r.get('Mapped Module ID') or None; mod = next((m for m in modules if m['catalogId'] == mo), None)
    text = r['Problem / exercise']; first = re.split(r'(?<=[.:?])\s', text, maxsplit=1)[0]
    base = {'id': qid, 'moduleId': mid(mo) if mo else None, 'module': mod['name'] if mod else None, 'topicId': tid, 'topic': topicById[tid]['name'] if tid in topicById else r['Topic'],
        'platform': r['Platform'], 'difficulty': r['Difficulty'], 'prerequisites': txt(r['Prerequisites']), 'url': txt(r['Direct URL']), 'how': txt(r['How to practise']),
        'workbookStatus': txt(r['Status']), 'workbookNotes': txt(r['Mistake / insight notes']), 'sessionDays': days, 'mappingConfidence': txt(r.get('Mapping confidence')),
        'workbookTopicId': r['Topic ID'], 'source': 'Workbook v7'}
    if qid in qmap:
        qmap[qid].update({k: v for k, v in base.items() if v not in (None, '') and k not in ('source', 'topic')})
        qmap[qid]['topic'] = base['topic'] if qid.startswith('P00') and len(qid) == 4 and False else qmap[qid].get('topic', base['topic'])
        updated += 1
    else:
        base.update({'title': first if len(first) <= 150 else first[:147].rstrip() + '…', 'type': 'Day practice', 'prompt': text, 'acceptanceCriteria': None, 'hint': None, 'estimatedMinutes': 30, 'tags': []})
        qmap[qid] = base; bank['questions'].append(base); added += 1
bank['meta'].update({'total': len(bank['questions']), 'workbookQuestionsImported': sum(1 for q in bank['questions'] if q['id'].startswith('P')), 'workbookQuestionsInWorkbook': 2089,
    'note': 'Workbook v7 lists 2,089 questions. Only those present in the supplied export are included; the rest are added when the full workbook is imported.'})

# ---------- topic project ladder (separate file, loaded only when needed)
ladder = [{'id': r['Project ID'], 'topicId': r['Topic ID'], 'moduleId': mid(r['Module ID']), 'size': r['Size'], 'title': r['Project title'], 'purpose': r['Problem / purpose'], 'scope': r['Core scope and deliverables'],
    'tech': r['Suggested technology'], 'prerequisites': r['Prerequisites / unlock condition'], 'acceptance': r['Acceptance / quality bar'], 'signal': r['Company / portfolio signal'],
    'firstDay': num(r['First learning day']), 'completionDay': num(r['Topic completion day']), 'status': r['Status']} for r in rows('Topic Project Ladder')]

# ---------- extras and review queue
extras = {}
for key, sheet in [('resources', 'Free Resources'), ('github', 'GitHub Publishing'), ('android', 'Android Side Track'), ('languages', 'Language Strategy'),
                   ('interview', 'Interview Preparation'), ('pause', 'Academic Pause Planner'), ('qa', 'Workbook QA')]:
    extras[key] = rows(sheet)
extras['startHere'] = [r for r in raw_rows('Start Here') if True] if 'Start Here' in S else []
fg = raw_rows('Folder Guidance') if 'Folder Guidance' in S else []
hi = next((i for i, r in enumerate(fg) if r and r[0] == 'Folder / path'), None)
extras['folders'] = [dict(zip(fg[hi], r + [''] * (len(fg[hi]) - len(r)))) for r in fg[hi + 1:]] if hi is not None else []
extras['reviewQueue'] = rows('Mapping Review Queue')
if old and not extras['startHere']: extras['startHere'] = old['extras'].get('startHere', [])

out = {'meta': {'source': os.path.basename(src), 'version': 'v7', 'total': total, 'groupSize': 30, 'revisionOffsets': [3, 7, 21, 45],
    'note': 'Revision offsets count curriculum sessions, not calendar days.', 'sessionsFromWorkbook': total - len(carried), 'carriedOverDays': carried,
    'topicCount': len(topics), 'ladderCount': len(ladder)},
    'modules': modules, 'topics': topics, 'sessions': sessions, 'projects': projects, 'projectTopics': projectTopics, 'practice': [], 'extras': extras}
json.dump(out, open(os.path.join(DATA, 'curriculum.json'), 'w'), ensure_ascii=False, separators=(',', ':'))
json.dump({'items': ladder}, open(os.path.join(DATA, 'topic-projects.json'), 'w'), ensure_ascii=False, separators=(',', ':'))
json.dump(bank, open(bankPath, 'w'), ensure_ascii=False, separators=(',', ':'))
print('sessions', len(sessions), 'carried-over', len(carried), 'modules', len(modules), 'topics', len(topics), 'projects', len(projects), 'versions', sum(len(p['versions']) for p in projects),
      'ladder', len(ladder), 'questions added', added, 'updated', updated, 'bank total', len(bank['questions']))
